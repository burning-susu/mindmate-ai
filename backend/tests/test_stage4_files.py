from __future__ import annotations

import io
import json
import os
import struct
import time
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import update

from mindmate.api import files as files_api
from mindmate.application import files as file_service
from mindmate.application import parse_worker_service
from mindmate.application.files import (
    FileValidationError,
    resolve_storage_path,
    validate_managed_content_path,
)
from mindmate.config import Settings
from mindmate.infrastructure.models import FileRecord, Folder
from mindmate.main import create_app


@pytest.fixture
def file_client(tmp_path: Path):
    settings = Settings(data_dir=tmp_path, env="test")
    with TestClient(create_app(settings), base_url="http://127.0.0.1") as client:
        response = client.post(
            "/api/v1/system/session", headers={"Origin": "http://127.0.0.1:5173"}
        )
        assert response.status_code == 200
        yield client, tmp_path


def write_headers(key: str) -> dict[str, str]:
    return {"Origin": "http://127.0.0.1:5173", "Idempotency-Key": key}


def import_one(client: TestClient, name: str, content: bytes, key: str):
    return client.post(
        "/api/v1/file-imports",
        headers=write_headers(key),
        files={"files": (name, content, "text/plain")},
    )


def wait_for_import(client: TestClient, import_id: str) -> dict:
    for _ in range(200):
        payload = client.get(f"/api/v1/file-imports/{import_id}").json()
        if payload["status"] in {"COMPLETED", "FAILED", "CANCELLED", "BLOCKED"}:
            return payload
        time.sleep(0.02)
    raise AssertionError(f"导入任务未在测试窗口内完成: {import_id}")


def wait_for_file(client: TestClient, file_id: str, statuses: set[str]) -> dict:
    for _ in range(200):
        payload = client.get(f"/api/v1/files/{file_id}").json()
        if payload.get("status") in statuses:
            return payload
        time.sleep(0.02)
    raise AssertionError(f"文件状态未在测试窗口内达到 {statuses}: {file_id}")


def create_folder(client: TestClient, name: str, key: str, parent_id: str | None = None):
    return client.post(
        "/api/v1/folders",
        headers=write_headers(key),
        json={"name": name, "parent_folder_id": parent_id},
    )


def test_text_import_is_streamed_into_controlled_storage(file_client) -> None:
    client, data_dir = file_client
    response = import_one(client, "notes.txt", "第一行\n第二行".encode(), "import-text-1")

    assert response.status_code == 202
    payload = response.json()
    assert payload["status"] in {"QUEUED", "COMPLETED"}
    item = payload["items"][0]
    assert item["status"] == "IMPORTED"
    file_id = item["file_id"]
    wait_for_import(client, payload["import_id"])

    detail = client.get(f"/api/v1/files/{file_id}")
    assert detail.status_code == 200
    assert detail.json()["status"] == "PARSED"
    assert detail.json()["content_available"] is True

    text = client.get(f"/api/v1/files/{file_id}/text")
    assert text.status_code == 200
    assert "第二行" in text.json()["text"]

    content = client.get(f"/api/v1/files/{file_id}/content")
    assert content.status_code == 200
    assert content.headers["x-content-type-options"] == "nosniff"
    assert content.content == "第一行\n第二行".encode()
    assert list((data_dir / "objects").rglob("*.txt"))
    assert not list((data_dir / "tasks").rglob("*.staging"))


def test_duplicate_requires_decision_and_separate_record_reuses_content(file_client) -> None:
    client, data_dir = file_client
    first = import_one(client, "same.txt", b"same-content", "import-duplicate-1")
    first_file_id = first.json()["items"][0]["file_id"]
    duplicate = import_one(client, "copy.txt", b"same-content", "import-duplicate-2")

    assert duplicate.status_code == 202
    duplicate_payload = duplicate.json()
    assert duplicate_payload["status"] == "BLOCKED"
    assert duplicate_payload["items"][0]["duplicate_status"] == "PENDING_DECISION"
    import_id = duplicate_payload["import_id"]
    decision = client.post(
        f"/api/v1/file-imports/{import_id}/duplicate-decisions",
        headers=write_headers("duplicate-decision-1"),
        json={"decisions": [{"item_index": 0, "decision": "CREATE_SEPARATE_RECORD"}]},
    )
    assert decision.status_code == 200
    second_file_id = decision.json()["items"][0]["file_id"]
    assert second_file_id != first_file_id
    assert len([path for path in (data_dir / "objects").rglob("*") if path.is_file()]) == 1

    listed = client.get("/api/v1/files").json()["items"]
    assert {item["file_id"] for item in listed} == {first_file_id, second_file_id}


def test_reusing_duplicate_moves_file_to_selected_folder(file_client) -> None:
    client, _ = file_client
    folder = create_folder(client, "面试题", "duplicate-folder-target").json()
    first = import_one(client, "interview.md", b"same-interview-content", "import-folder-duplicate-1")
    first_file_id = first.json()["items"][0]["file_id"]
    duplicate = client.post(
        "/api/v1/file-imports",
        headers=write_headers("import-folder-duplicate-2"),
        data={"folder_id": folder["folder_id"]},
        files={"files": ("interview-copy.md", b"same-interview-content", "text/markdown")},
    )
    assert duplicate.status_code == 202
    import_id = duplicate.json()["import_id"]

    decided = client.post(
        f"/api/v1/file-imports/{import_id}/duplicate-decisions",
        headers=write_headers("reuse-into-folder"),
        json={"decisions": [{"item_index": 0, "decision": "REUSE_EXISTING"}]},
    )
    assert decided.status_code == 200
    assert decided.json()["items"][0]["file_id"] == first_file_id

    file_detail = client.get(f"/api/v1/files/{first_file_id}").json()
    assert file_detail["folder_id"] == folder["folder_id"]
    assert client.get(f"/api/v1/files?folder_id={folder['folder_id']}").json()["items"][0]["file_id"] == first_file_id
    folder_list = client.get("/api/v1/folders").json()["items"]
    assert next(item["file_count"] for item in folder_list if item["folder_id"] == folder["folder_id"]) == 1


def test_reusing_duplicate_restores_matching_file_from_trash(file_client) -> None:
    client, _ = file_client
    first = import_one(client, "trashed.txt", b"trashed-content", "import-trashed-duplicate-1")
    first_file_id = first.json()["items"][0]["file_id"]
    wait_for_file(client, first_file_id, {"PARSED"})
    first_detail = client.get(f"/api/v1/files/{first_file_id}").json()
    trashed = client.delete(
        f"/api/v1/files/{first_file_id}",
        params={"expected_version": first_detail["row_version"]},
        headers=write_headers("trash-duplicate-source"),
    )
    assert trashed.status_code == 200

    duplicate = import_one(client, "copy.txt", b"trashed-content", "import-trashed-duplicate-2")
    duplicate_payload = duplicate.json()
    assert duplicate_payload["status"] == "BLOCKED"

    decided = client.post(
        f"/api/v1/file-imports/{duplicate_payload['import_id']}/duplicate-decisions",
        headers=write_headers("reuse-trashed-duplicate"),
        json={"decisions": [{"item_index": 0, "decision": "REUSE_EXISTING"}]},
    )
    assert decided.status_code == 200
    assert decided.json()["status"] == "COMPLETED"
    assert decided.json()["items"][0]["status"] == "REUSED"
    assert decided.json()["items"][0]["file_id"] == first_file_id

    restored = client.get(f"/api/v1/files/{first_file_id}").json()
    assert restored["deleted_at"] is None
    assert restored["status"] == "PARSED"
    assert [item["file_id"] for item in client.get("/api/v1/files").json()["items"]] == [first_file_id]
    assert client.get("/api/v1/trash").json()["files"] == []


def test_knowledge_base_import_reuses_source_and_persists_membership(file_client) -> None:
    client, _ = file_client
    knowledge_base = client.post(
        "/api/v1/knowledge-bases",
        headers=write_headers("import-context-kb"),
        json={"name": "导入上下文库", "color": "#176b87"},
    )
    assert knowledge_base.status_code == 201
    knowledge_base_id = knowledge_base.json()["knowledge_base_id"]

    first = import_one(client, "source.txt", b"shared-source", "import-context-source")
    first_file_id = first.json()["items"][0]["file_id"]
    duplicate = client.post(
        "/api/v1/file-imports",
        headers=write_headers("import-context-duplicate"),
        data={"knowledge_base_id": knowledge_base_id},
        files={"files": ("copy.txt", b"shared-source", "text/plain")},
    )
    assert duplicate.status_code == 202
    duplicate_payload = duplicate.json()
    assert duplicate_payload["status"] == "BLOCKED"

    decided = client.post(
        f"/api/v1/file-imports/{duplicate_payload['import_id']}/duplicate-decisions",
        headers=write_headers("import-context-reuse"),
        json={"decisions": [{"item_index": 0, "decision": "REUSE_EXISTING"}]},
    )
    assert decided.status_code == 200
    assert decided.json()["items"][0]["file_id"] == first_file_id
    refreshed_decided = client.get(f"/api/v1/file-imports/{duplicate_payload['import_id']}")
    assert refreshed_decided.status_code == 200
    assert refreshed_decided.json()["status"] == "COMPLETED"
    assert refreshed_decided.json()["items"][0]["status"] == "REUSED"
    assert refreshed_decided.json()["items"][0]["duplicate_status"] == "REUSED"
    assert refreshed_decided.json()["items"][0]["file_id"] == first_file_id

    members = client.get(f"/api/v1/knowledge-bases/{knowledge_base_id}/files")
    assert members.status_code == 200
    assert [item["file_id"] for item in members.json()["items"]] == [first_file_id]
    assert client.get(f"/api/v1/knowledge-bases/{knowledge_base_id}").json()["status"] == "PREPARING"

    direct = client.post(
        "/api/v1/file-imports",
        headers=write_headers("import-context-direct"),
        data={"knowledge_base_id": knowledge_base_id},
        files={"files": ("new.txt", b"new-source", "text/plain")},
    )
    assert direct.status_code == 202
    direct_file_id = direct.json()["items"][0]["file_id"]
    assert direct_file_id != first_file_id
    assert {item["file_id"] for item in client.get(f"/api/v1/knowledge-bases/{knowledge_base_id}/files").json()["items"]} == {first_file_id, direct_file_id}

    skipped = client.post(
        "/api/v1/file-imports",
        headers=write_headers("import-context-skip"),
        data={"knowledge_base_id": knowledge_base_id},
        files={"files": ("skip.txt", b"shared-source", "text/plain")},
    )
    assert skipped.status_code == 202
    skipped_payload = skipped.json()
    assert skipped_payload["status"] == "BLOCKED"
    skipped_decision = client.post(
        f"/api/v1/file-imports/{skipped_payload['import_id']}/duplicate-decisions",
        headers=write_headers("import-context-skip-decision"),
        json={"decisions": [{"item_index": 0, "decision": "SKIP"}]},
    )
    assert skipped_decision.status_code == 200
    refreshed_skipped = client.get(f"/api/v1/file-imports/{skipped_payload['import_id']}")
    assert refreshed_skipped.status_code == 200
    assert refreshed_skipped.json()["status"] == "COMPLETED"
    assert refreshed_skipped.json()["items"][0]["status"] == "SKIPPED"
    assert refreshed_skipped.json()["items"][0]["duplicate_status"] == "SKIPPED"
    assert refreshed_skipped.json()["items"][0]["file_id"] is None
    assert {item["file_id"] for item in client.get(f"/api/v1/knowledge-bases/{knowledge_base_id}/files").json()["items"]} == {first_file_id, direct_file_id}


def test_invalid_signature_and_batch_limits_are_rejected_per_item(file_client) -> None:
    client, _ = file_client
    disguised = client.post(
        "/api/v1/file-imports",
        headers=write_headers("import-invalid-signature"),
        files={"files": ("fake.pdf", b"plain text", "application/pdf")},
    )
    assert disguised.status_code == 202
    assert disguised.json()["items"][0]["error_code"] == "FILE_SIGNATURE_INVALID"

    too_many = client.post(
        "/api/v1/file-imports",
        headers=write_headers("import-too-many"),
        files=[("files", (f"{index}.txt", b"x", "text/plain")) for index in range(21)],
    )
    assert too_many.status_code == 202
    assert too_many.json()["items"][-1]["error_code"] == "TOO_MANY_FILES"


def test_folder_tag_trash_restore_and_purge(file_client) -> None:
    client, data_dir = file_client
    folder = client.post(
        "/api/v1/folders",
        headers=write_headers("folder-create-1"),
        json={"name": "资料"},
    )
    assert folder.status_code == 201
    folder_id = folder.json()["folder_id"]
    tag = client.post(
        "/api/v1/tags",
        headers=write_headers("tag-create-1"),
        json={"name": "重要", "color": "#176b87"},
    )
    assert tag.status_code == 201
    tag_id = tag.json()["tag_id"]
    imported = client.post(
        "/api/v1/file-imports",
        headers=write_headers("import-folder-tag-1"),
        data={"folder_id": folder_id, "tag_ids": f'["{tag_id}"]'},
        files={"files": ("organized.md", b"# title", "text/markdown")},
    )
    file_id = imported.json()["items"][0]["file_id"]
    assert imported.status_code == 202
    assert client.get(f"/api/v1/files/{file_id}").json()["folder_id"] == folder_id
    assert client.get("/api/v1/files?q=重要").json()["items"][0]["file_id"] == file_id

    wait_for_file(client, file_id, {"PARSED"})
    current_version = client.get(f"/api/v1/files/{file_id}").json()["row_version"]
    trashed = client.delete(
        f"/api/v1/files/{file_id}",
        params={"expected_version": current_version},
        headers=write_headers("trash-file-1"),
    )
    assert trashed.status_code == 200
    assert client.get("/api/v1/files").json()["items"] == []
    assert client.get("/api/v1/trash").json()["files"][0]["file_id"] == file_id

    restored = client.post(
        f"/api/v1/trash/file/{file_id}/restore",
        params={"expected_version": trashed.json()["row_version"]},
        headers=write_headers("restore-file-1"),
    )
    assert restored.status_code == 200
    assert client.get("/api/v1/files").json()["items"][0]["file_id"] == file_id

    deleted = client.delete(
        f"/api/v1/files/{file_id}",
        params={"expected_version": restored.json()["row_version"]},
        headers=write_headers("trash-file-2"),
    )
    assert deleted.status_code == 200
    purged = client.delete(
        f"/api/v1/trash/file/{file_id}",
        params={"confirmed": "true", "expected_version": deleted.json()["row_version"]},
        headers=write_headers("purge-file-1"),
    )
    assert purged.status_code == 200
    assert client.get(f"/api/v1/files/{file_id}").status_code == 404
    assert not [path for path in (data_dir / "objects").rglob("*") if path.is_file()]


def test_folder_tree_restore_and_permanent_purge(file_client) -> None:
    client, data_dir = file_client
    root = client.post(
        "/api/v1/folders", headers=write_headers("tree-root-1"), json={"name": "根"}
    ).json()
    child = client.post(
        "/api/v1/folders",
        headers=write_headers("tree-child-1"),
        json={"name": "子", "parent_folder_id": root["folder_id"]},
    ).json()
    imported = client.post(
        "/api/v1/file-imports",
        headers=write_headers("tree-import-1"),
        data={"folder_id": child["folder_id"]},
        files={"files": ("tree.txt", b"tree", "text/plain")},
    ).json()
    file_id = imported["items"][0]["file_id"]

    deleted = client.delete(
        f"/api/v1/folders/{root['folder_id']}",
        params={"deletion_strategy": "TRASH_RECURSIVE", "expected_version": 1},
        headers=write_headers("tree-trash-1"),
    )
    assert deleted.status_code == 200
    assert client.get("/api/v1/files").json()["items"] == []
    restored = client.post(
        f"/api/v1/trash/folder/{root['folder_id']}/restore",
        params={"expected_version": deleted.json()["row_version"]},
        headers=write_headers("tree-restore-1"),
    )
    assert restored.status_code == 200
    assert client.get("/api/v1/files").json()["items"][0]["file_id"] == file_id

    deleted_again = client.delete(
        f"/api/v1/folders/{root['folder_id']}",
        params={
            "deletion_strategy": "TRASH_RECURSIVE",
            "expected_version": restored.json()["row_version"],
        },
        headers=write_headers("tree-trash-2"),
    )
    purged = client.delete(
        f"/api/v1/trash/folder/{root['folder_id']}",
        params={"confirmed": "true", "expected_version": deleted_again.json()["row_version"]},
        headers=write_headers("tree-purge-1"),
    )
    assert purged.status_code == 200
    assert client.get(f"/api/v1/files/{file_id}").status_code == 404
    assert not [path for path in (data_dir / "objects").rglob("*") if path.is_file()]


def test_office_documents_use_isolated_parsers_and_blank_pdf_is_not_marked_parsed(
    file_client, monkeypatch
) -> None:
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    from pypdf import PdfWriter

    client, data_dir = file_client
    docx_buffer = io.BytesIO()
    document = Document()
    document.add_heading("学习主题", level=1)
    document.add_paragraph("这是 Word 正文。")
    document.save(docx_buffer)
    docx_response = client.post(
        "/api/v1/file-imports",
        headers=write_headers("import-docx-1"),
        files={
            "files": (
                "lesson.docx",
                docx_buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    docx_item = docx_response.json()["items"][0]
    assert docx_item["parse_status"] in {"QUEUED", "PARSED"}
    wait_for_file(client, docx_item["file_id"], {"PARSED"})
    docx_preview = client.get(f"/api/v1/files/{docx_item['file_id']}/preview").json()
    assert "这是 Word 正文" in docx_preview["text"]
    assert docx_preview["metadata"]["paragraph_count"] >= 2
    rendered_docx_pdf = data_dir / "rendered-docx-preview.pdf"
    docx_writer = PdfWriter()
    docx_writer.add_blank_page(width=72, height=72)
    docx_writer.write(str(rendered_docx_pdf))
    monkeypatch.setattr(
        files_api,
        "get_or_create_office_preview",
        lambda _source, _sha, document_type, _settings: (rendered_docx_pdf, 1, False)
        if document_type == "DOCX"
        else (rendered_docx_pdf, 1, False),
    )
    rendered_docx_response = client.get(f"/api/v1/files/{docx_item['file_id']}/docx-preview")
    assert rendered_docx_response.status_code == 200
    assert rendered_docx_response.headers["content-type"] == "application/pdf"
    assert rendered_docx_response.content.startswith(b"%PDF-")

    pptx_buffer = io.BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    text_box.text = "第一张幻灯片"
    presentation.save(pptx_buffer)
    pptx_response = client.post(
        "/api/v1/file-imports",
        headers=write_headers("import-pptx-1"),
        files={
            "files": (
                "slides.pptx",
                pptx_buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        },
    )
    pptx_item = pptx_response.json()["items"][0]
    assert pptx_item["parse_status"] in {"QUEUED", "PARSED"}
    wait_for_file(client, pptx_item["file_id"], {"PARSED"})
    pptx_preview = client.get(f"/api/v1/files/{pptx_item['file_id']}/preview").json()
    assert "第一张幻灯片" in pptx_preview["text"]
    assert pptx_preview["metadata"]["slide_count"] == 1
    rendered_pdf = data_dir / "rendered-preview.pdf"
    rendered = PdfWriter()
    rendered.add_blank_page(width=72, height=72)
    rendered.write(str(rendered_pdf))
    monkeypatch.setattr(
        files_api,
        "get_or_create_office_preview",
        lambda _source, _sha, document_type, _settings: (rendered_pdf, 1, False),
    )
    rendered_response = client.get(f"/api/v1/files/{pptx_item['file_id']}/pptx-preview")
    assert rendered_response.status_code == 200
    assert rendered_response.headers["content-type"] == "application/pdf"
    assert rendered_response.headers["x-content-type-options"] == "nosniff"
    assert rendered_response.content.startswith(b"%PDF-")

    monkeypatch.setattr(
        files_api,
        "get_or_create_office_preview",
        lambda _source, _sha, _document_type, _settings: (rendered_pdf, 2, False),
    )
    mismatch_response = client.get(f"/api/v1/files/{pptx_item['file_id']}/pptx-preview")
    assert mismatch_response.status_code == 422
    assert mismatch_response.json()["code"] == "PPTX_RENDER_PAGE_COUNT_MISMATCH"

    monkeypatch.undo()
    from mindmate.application import pptx_preview as pptx_preview_service

    monkeypatch.setattr(pptx_preview_service, "find_libreoffice", lambda _settings: None)
    unavailable_response = client.get(f"/api/v1/files/{pptx_item['file_id']}/pptx-preview")
    assert unavailable_response.status_code == 503
    assert unavailable_response.json()["code"] == "OFFICE_RENDERER_UNAVAILABLE"
    unavailable_docx_response = client.get(f"/api/v1/files/{docx_item['file_id']}/docx-preview")
    assert unavailable_docx_response.status_code == 503
    assert unavailable_docx_response.json()["code"] == "OFFICE_RENDERER_UNAVAILABLE"

    pdf_buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(pdf_buffer)
    pdf_response = client.post(
        "/api/v1/file-imports",
        headers=write_headers("import-pdf-blank-1"),
        files={"files": ("scan.pdf", pdf_buffer.getvalue(), "application/pdf")},
    )
    pdf_item = pdf_response.json()["items"][0]
    assert pdf_item["status"] == "IMPORTED"
    pdf_task = wait_for_import(client, pdf_response.json()["import_id"])
    pdf_result = pdf_task["items"][0]
    assert pdf_result["parse_status"] == "PARSE_FAILED"
    assert "OCR" in pdf_result["parse_error"]


def test_xmind_json_and_xml_import_as_searchable_topic_trees(file_client) -> None:
    client, _ = file_client
    modern = [
        {
            "class": "sheet",
            "title": "职业规划",
            "rootTopic": {
                "id": "root",
                "title": "前端工程师",
                "children": {
                    "attached": [
                        {
                            "id": "html",
                            "title": "HTML 基础",
                            "children": {"attached": [{"id": "semantic", "title": "语义化标签"}]},
                        },
                        {"id": "css", "title": "CSS 布局"},
                    ]
                },
            },
        }
    ]
    modern_buffer = io.BytesIO()
    native_thumbnail = bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
        "0000000b49444154789c636000020000050001a5f645400000000049454e44ae426082"
    )
    with zipfile.ZipFile(modern_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("content.json", json.dumps(modern, ensure_ascii=False))
        archive.writestr("metadata.json", "{}")
        archive.writestr("Thumbnails/thumbnail.png", native_thumbnail)
    modern_response = client.post(
        "/api/v1/file-imports",
        headers=write_headers("import-xmind-json"),
        files={"files": ("职业规划.xmind", modern_buffer.getvalue(), "application/vnd.xmind.workbook")},
    )
    assert modern_response.status_code == 202
    modern_item = modern_response.json()["items"][0]
    assert modern_item["status"] == "IMPORTED"
    wait_for_file(client, modern_item["file_id"], {"PARSED"})
    modern_preview = client.get(f"/api/v1/files/{modern_item['file_id']}/preview").json()
    assert "语义化标签" in modern_preview["text"]
    assert modern_preview["metadata"]["sheet_count"] == 1
    assert modern_preview["metadata"]["topic_count"] == 4
    outline = modern_preview["metadata"]["xmind_outline"][0]
    assert outline["title"] == "职业规划"
    assert outline["root"]["children"][0]["children"][0]["title"] == "语义化标签"
    thumbnail_response = client.get(f"/api/v1/files/{modern_item['file_id']}/xmind-thumbnail")
    assert thumbnail_response.status_code == 200
    assert thumbnail_response.headers["content-type"] == "image/png"
    assert thumbnail_response.headers["x-content-type-options"] == "nosniff"
    assert thumbnail_response.content == native_thumbnail

    classic_xml = '''<?xml version="1.0" encoding="UTF-8"?>
    <xmap-content xmlns="urn:xmind:xmap:xmlns:content:2.0" version="2.0">
      <sheet id="s1"><title>旧版工作表</title>
        <topic id="r"><title>核心主题</title><children><topics type="attached">
          <topic id="c"><title>子主题</title></topic>
        </topics></children></topic>
      </sheet>
    </xmap-content>'''.encode()
    classic_buffer = io.BytesIO()
    with zipfile.ZipFile(classic_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("content.xml", classic_xml)
        archive.writestr("META-INF/manifest.xml", "<manifest/>")
    classic_response = client.post(
        "/api/v1/file-imports",
        headers=write_headers("import-xmind-xml"),
        files={"files": ("classic.xmind", classic_buffer.getvalue(), "application/x-xmind")},
    )
    assert classic_response.status_code == 202
    classic_item = classic_response.json()["items"][0]
    assert classic_item["status"] == "IMPORTED"
    wait_for_file(client, classic_item["file_id"], {"PARSED"})
    classic_preview = client.get(f"/api/v1/files/{classic_item['file_id']}/preview").json()
    assert "子主题" in classic_preview["text"]
    assert classic_preview["metadata"]["xmind_outline"][0]["root"]["children"][0]["title"] == "子主题"
    assert client.get(f"/api/v1/files/{classic_item['file_id']}/xmind-thumbnail").status_code == 204


def test_xmind_extension_requires_an_xmind_zip_signature(file_client) -> None:
    client, _ = file_client
    fake = io.BytesIO()
    with zipfile.ZipFile(fake, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("notes.txt", "not an XMind map")
    response = client.post(
        "/api/v1/file-imports",
        headers=write_headers("fake-xmind-package"),
        files={"files": ("fake.xmind", fake.getvalue(), "application/zip")},
    )
    assert response.status_code == 202
    assert response.json()["items"][0]["error_code"] == "FILE_SIGNATURE_INVALID"


def test_xmind_thumbnail_rejects_excessive_pixel_dimensions(tmp_path: Path) -> None:
    image = bytearray(24)
    image[:8] = b"\x89PNG\r\n\x1a\n"
    image[12:16] = b"IHDR"
    image[16:24] = struct.pack(">II", 10_000, 10_000)
    package = tmp_path / "large-thumbnail.xmind"
    with zipfile.ZipFile(package, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("Thumbnails/thumbnail.png", image)

    assert files_api._read_xmind_thumbnail(package) is None


def test_mime_size_and_batch_limits_are_enforced(file_client, monkeypatch) -> None:
    client, _ = file_client
    mismatch = client.post(
        "/api/v1/file-imports",
        headers=write_headers("mime-mismatch-1"),
        files={"files": ("fake.txt", b"plain", "application/pdf")},
    )
    assert mismatch.json()["items"][0]["error_code"] == "FILE_MIME_INVALID"

    monkeypatch.setattr(file_service, "MAX_FILE_SIZE", 4)
    size_response = import_one(client, "large.txt", b"12345", "size-limit-1")
    assert size_response.json()["items"][0]["error_code"] == "FILE_TOO_LARGE"

    monkeypatch.setattr(file_service, "MAX_FILE_SIZE", 50 * 1024 * 1024)
    monkeypatch.setattr(files_api, "MAX_BATCH_SIZE", 8)
    batch_response = client.post(
        "/api/v1/file-imports",
        headers=write_headers("batch-limit-1"),
        files=[
            ("files", ("one.txt", b"12345", "text/plain")),
            ("files", ("two.txt", b"67890", "text/plain")),
        ],
    )
    assert batch_response.json()["items"][1]["error_code"] == "BATCH_TOO_LARGE"


def test_storage_paths_reject_traversal_absolute_and_symlink_escape(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path / "data", env="test")
    settings.ensure_data_dirs()
    with pytest.raises(FileValidationError, match="受控文件路径无效"):
        resolve_storage_path(settings, "../../outside.txt")
    with pytest.raises(FileValidationError, match="受控文件路径无效"):
        resolve_storage_path(settings, str((tmp_path / "outside.txt").resolve()))

    outside = tmp_path / "outside"
    outside.mkdir()
    link = settings.resolved_data_dir / "objects" / "escape"
    symlink_created = True
    try:
        os.symlink(outside, link, target_is_directory=True)
    except OSError:
        symlink_created = False
    if symlink_created:
        with pytest.raises(FileValidationError, match="受控文件路径"):
            validate_managed_content_path(settings, "objects/escape/secret.txt")

    outside_file = outside / "outside.txt"
    outside_file.write_text("secret", encoding="utf-8")
    hardlink = settings.resolved_data_dir / "objects" / "hardlink.txt"
    os.link(outside_file, hardlink)
    with pytest.raises(FileValidationError, match="硬链接"):
        validate_managed_content_path(settings, "objects/hardlink.txt")


def test_folder_guards_null_semantics_and_multi_level_trash(file_client) -> None:
    client, _ = file_client
    missing_parent = create_folder(client, "orphan", "folder-missing-parent", "missing")
    assert missing_parent.status_code == 404

    root = create_folder(client, "root", "folder-root-guard").json()
    child = create_folder(client, "child", "folder-child-guard", root["folder_id"]).json()
    grandchild = create_folder(
        client, "grandchild", "folder-grandchild-guard", child["folder_id"]
    ).json()
    self_move = client.post(
        f"/api/v1/folders/{root['folder_id']}/move",
        headers=write_headers("folder-self-move"),
        json={"parent_folder_id": root["folder_id"], "row_version": 1},
    )
    assert self_move.status_code == 400
    cycle_move = client.post(
        f"/api/v1/folders/{root['folder_id']}/move",
        headers=write_headers("folder-cycle-move"),
        json={"parent_folder_id": grandchild["folder_id"], "row_version": 1},
    )
    assert cycle_move.status_code == 400

    tag = client.post(
        "/api/v1/tags",
        headers=write_headers("tag-create-color"),
        json={"name": "重点", "color": "#ff0000"},
    ).json()
    cleared = client.patch(
        f"/api/v1/tags/{tag['tag_id']}",
        headers=write_headers("tag-clear-color"),
        json={"color": None, "row_version": tag["row_version"]},
    )
    assert cleared.json()["color"] is None

    imported = client.post(
        "/api/v1/file-imports",
        headers=write_headers("nested-file-import"),
        data={"folder_id": grandchild["folder_id"]},
        files={"files": ("nested.txt", b"nested", "text/plain")},
    ).json()
    file_id = imported["items"][0]["file_id"]
    wait_for_file(client, file_id, {"PARSED"})
    deleted = client.delete(
        f"/api/v1/folders/{root['folder_id']}?deletion_strategy=TRASH_RECURSIVE&expected_version={root['row_version']}",
        headers=write_headers("folder-trash-deep"),
    )
    restored = client.post(
        f"/api/v1/trash/folder/{root['folder_id']}/restore",
        params={"expected_version": deleted.json()["row_version"]},
        headers=write_headers("folder-restore-deep"),
    )
    assert restored.status_code == 200
    assert client.get(f"/api/v1/files/{file_id}").json()["status"] == "PARSED"

    deleted_again = client.delete(
        f"/api/v1/folders/{root['folder_id']}?deletion_strategy=TRASH_RECURSIVE&expected_version={restored.json()['row_version']}",
        headers=write_headers("folder-trash-deep-2"),
    )
    purged = client.delete(
        f"/api/v1/trash/folder/{root['folder_id']}?confirmed=true&expected_version={deleted_again.json()['row_version']}",
        headers=write_headers("folder-purge-deep"),
    )
    assert purged.status_code == 200
    assert client.get(f"/api/v1/files/{file_id}").status_code == 404


def test_corrupt_folder_cycle_is_rejected_without_hanging(file_client) -> None:
    client, _ = file_client
    first = create_folder(client, "first", "cycle-first").json()
    second = create_folder(client, "second", "cycle-second", first["folder_id"]).json()
    factory = client.app.state.session_factory
    with factory() as session:
        session.execute(
            update(Folder)
            .where(Folder.folder_id == first["folder_id"])
            .values(parent_folder_id=second["folder_id"])
        )
        session.commit()
    response = client.post(
        f"/api/v1/folders/{first['folder_id']}/move",
        headers=write_headers("cycle-detected"),
        json={"parent_folder_id": None, "row_version": 1},
    )
    assert response.status_code == 409
    assert response.json()["code"] == "FOLDER_TREE_INVALID"


def test_folder_and_tag_optimistic_locks_are_atomic(file_client) -> None:
    client, _ = file_client
    folder = create_folder(client, "资料", "lock-folder-create").json()
    target = create_folder(client, "归档", "lock-target-create").json()
    assert folder["row_version"] == 1

    renamed = client.patch(
        f"/api/v1/folders/{folder['folder_id']}",
        headers=write_headers("lock-folder-rename"),
        json={"name": "新资料", "row_version": 1},
    )
    assert renamed.status_code == 200
    assert renamed.json()["row_version"] == 2

    stale_rename = client.patch(
        f"/api/v1/folders/{folder['folder_id']}",
        headers=write_headers("lock-folder-stale-rename"),
        json={"name": "旧客户端覆盖", "row_version": 1},
    )
    assert stale_rename.status_code == 412
    assert stale_rename.json()["code"] == "RESOURCE_VERSION_CONFLICT"
    assert stale_rename.json()["current_row_version"] == 2
    assert client.get(f"/api/v1/folders/{folder['folder_id']}").json()["name"] == "新资料"

    moved = client.post(
        f"/api/v1/folders/{folder['folder_id']}/move",
        headers=write_headers("lock-folder-move"),
        json={"parent_folder_id": target["folder_id"], "row_version": 2},
    )
    assert moved.status_code == 200
    assert moved.json()["row_version"] == 3

    stale_delete = client.delete(
        f"/api/v1/folders/{folder['folder_id']}",
        params={"deletion_strategy": "TRASH_RECURSIVE", "expected_version": 2},
        headers=write_headers("lock-folder-stale-delete"),
    )
    assert stale_delete.status_code == 412
    assert client.get(f"/api/v1/folders/{folder['folder_id']}").status_code == 200

    deleted = client.delete(
        f"/api/v1/folders/{folder['folder_id']}",
        params={"deletion_strategy": "TRASH_RECURSIVE", "expected_version": 3},
        headers=write_headers("lock-folder-delete"),
    )
    assert deleted.status_code == 200
    assert deleted.json()["row_version"] == 4

    stale_restore = client.post(
        f"/api/v1/trash/folder/{folder['folder_id']}/restore",
        params={"expected_version": 3},
        headers=write_headers("lock-folder-stale-restore"),
    )
    assert stale_restore.status_code == 412
    restored = client.post(
        f"/api/v1/trash/folder/{folder['folder_id']}/restore",
        params={"expected_version": 4},
        headers=write_headers("lock-folder-restore"),
    )
    assert restored.status_code == 200
    assert restored.json()["row_version"] == 5

    tag = client.post(
        "/api/v1/tags",
        headers=write_headers("lock-tag-create"),
        json={"name": "重点", "color": "#176b87"},
    ).json()
    assert tag["row_version"] == 1
    updated_tag = client.patch(
        f"/api/v1/tags/{tag['tag_id']}",
        headers=write_headers("lock-tag-update"),
        json={"color": "#123456", "row_version": 1},
    )
    assert updated_tag.status_code == 200
    assert updated_tag.json()["row_version"] == 2

    stale_tag = client.patch(
        f"/api/v1/tags/{tag['tag_id']}",
        headers=write_headers("lock-tag-stale"),
        json={"color": "#ffffff", "row_version": 1},
    )
    assert stale_tag.status_code == 412
    current_tag = next(
        item for item in client.get("/api/v1/tags").json()["items"] if item["tag_id"] == tag["tag_id"]
    )
    assert current_tag["color"] == "#123456"

    stale_tag_delete = client.delete(
        f"/api/v1/tags/{tag['tag_id']}",
        params={"expected_version": 1},
        headers=write_headers("lock-tag-stale-delete"),
    )
    assert stale_tag_delete.status_code == 412
    deleted_tag = client.delete(
        f"/api/v1/tags/{tag['tag_id']}",
        params={"expected_version": 2},
        headers=write_headers("lock-tag-delete"),
    )
    assert deleted_tag.status_code == 200
    missing_tag = client.delete(
        f"/api/v1/tags/{tag['tag_id']}",
        params={"expected_version": 2},
        headers=write_headers("lock-tag-delete-missing"),
    )
    assert missing_tag.status_code == 404


def test_parse_failure_fields_persist_and_successful_retry_clears_error(
    file_client, monkeypatch
) -> None:
    client, data_dir = file_client
    private_path = str(data_dir / "private" / "notes.txt")

    def fail_parse(_path: Path):
        raise FileValidationError(
            "PARSER_FAILED", f"Traceback: failed while reading {private_path}"
        )

    client.app.state.settings.parse_worker_max_retries = 0
    monkeypatch.setattr(parse_worker_service, "parse_local_text", fail_parse)
    imported = import_one(client, "retry.txt", b"retry me", "parse-failure-persist")
    assert imported.status_code == 202
    item = imported.json()["items"][0]
    failed_import = wait_for_import(client, imported.json()["import_id"])
    item = failed_import["items"][0]
    assert item["parse_status"] == "PARSE_FAILED"
    assert item["parse_failure_stage"] == "PARSING"
    assert item["parse_error_id"] == "PARSER_FAILED"
    assert item["parse_retry_count"] == 0
    assert private_path not in item["parse_error"]
    assert "Traceback" not in item["parse_error"]

    file_id = item["file_id"]
    detail = client.get(f"/api/v1/files/{file_id}")
    assert detail.status_code == 200
    assert detail.json()["parse_failure_stage"] == "PARSING"
    assert detail.json()["parse_error_id"] == "PARSER_FAILED"
    assert detail.json()["parse_retry_count"] == 0
    assert private_path not in detail.text
    assert "Traceback" not in detail.text

    factory = client.app.state.session_factory
    with factory() as session:
        record = session.get(FileRecord, file_id)
        assert record is not None
        assert record.parse_failure_stage == "PARSING"
        assert record.parse_error_id == "PARSER_FAILED"
        assert record.parse_retry_count == 0

    monkeypatch.setattr(parse_worker_service, "parse_local_text", lambda _path: ("retry ok", {"line_count": 1}))
    retried = client.post(
        f"/api/v1/files/{file_id}/reprocess",
        headers=write_headers("parse-failure-retry"),
    )
    assert retried.status_code == 202
    retried_file = retried.json()["file"]
    wait_for_file(client, file_id, {"PARSED"})
    retried_file = client.get(f"/api/v1/files/{file_id}").json()
    assert retried_file["status"] == "PARSED"
    assert retried_file["parse_failure_stage"] is None
    assert retried_file["parse_error_id"] is None
    assert retried_file["parse_retry_count"] == 1
