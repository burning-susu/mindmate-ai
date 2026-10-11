from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any

MAX_PAGES = 2_000
MAX_TEXT_CHARACTERS = 10_000_000
MAX_XMIND_DEPTH = 100
MAX_XMIND_TOPICS = 200_000
MAX_ARCHIVE_ENTRIES = 10_000
MAX_ARCHIVE_EXPANDED_SIZE = 200 * 1024 * 1024
MAX_ARCHIVE_COMPRESSION_RATIO = 100
MAX_XMIND_CONTENT_SIZE = 20 * 1024 * 1024
MAX_SPREADSHEET_CELLS = 500_000
MAX_SPREADSHEET_SHARED_STRINGS = 500_000


class ParseFailure(Exception):
    pass


def _append(parts: list[str], locations: list[dict[str, Any]], text: str, **location: Any) -> None:
    value = text.strip()
    if not value:
        return
    current_size = int(locations[-1]["start"]) + int(locations[-1]["length"]) if locations else 0
    if current_size + len(value) > MAX_TEXT_CHARACTERS:
        raise ParseFailure("解析文本超过安全输出上限。")
    start = current_size
    parts.append(value)
    locations.append({"start": start, "length": len(value), **location})


def _parse_pdf(path: Path) -> dict[str, Any]:
    from pypdf import PdfReader

    reader = PdfReader(path, strict=True)
    if reader.is_encrypted:
        raise ParseFailure("加密 PDF 不受支持。")
    if len(reader.pages) > MAX_PAGES:
        raise ParseFailure("PDF 页数超过安全上限。")
    parts: list[str] = []
    locations: list[dict[str, Any]] = []
    for page_number, page in enumerate(reader.pages, start=1):
        _append(parts, locations, page.extract_text() or "", page=page_number)
    if not parts:
        raise ParseFailure("未检测到可解析文本；V1 不支持扫描 PDF OCR。")
    return {
        "text": "\n".join(parts),
        "metadata": {"page_count": len(reader.pages), "character_count": sum(len(p) for p in parts)},
        "locations": locations,
    }


def _parse_docx(path: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(str(path))
    parts: list[str] = []
    locations: list[dict[str, Any]] = []
    heading_path: list[str] = []
    paragraph_number = 0
    for paragraph in document.paragraphs:
        paragraph_number += 1
        text = paragraph.text.strip()
        style_name = (paragraph.style.name or "") if paragraph.style is not None else ""
        if style_name.startswith("Heading") and text:
            level_text = style_name.removeprefix("Heading").strip()
            level = int(level_text) if level_text.isdigit() else 1
            heading_path = heading_path[: max(0, level - 1)]
            heading_path.append(text)
            block_type = "TITLE"
        elif style_name.casefold().startswith("list"):
            block_type = "LIST"
        elif "code" in style_name.casefold():
            block_type = "CODE"
        else:
            block_type = "PARAGRAPH"
        _append(
            parts,
            locations,
            text,
            paragraph=paragraph_number,
            heading_path=list(heading_path),
            block_type=block_type,
        )
    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            _append(
                parts,
                locations,
                "\t".join(cell.text.strip() for cell in row.cells),
                table=table_number,
                row=row_number,
            )
    return {
        "text": "\n".join(parts),
        "metadata": {
            "paragraph_count": paragraph_number,
            "table_count": len(document.tables),
            "character_count": sum(len(p) for p in parts),
        },
        "locations": locations,
    }


def _shape_text(shape: Any) -> list[str]:
    values: list[str] = []
    if getattr(shape, "has_text_frame", False):
        value = str(shape.text).strip()
        if value:
            values.append(value)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            value = "\t".join(cell.text.strip() for cell in row.cells)
            if value.strip():
                values.append(value)
    return values


def _parse_pptx(path: Path) -> dict[str, Any]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    if len(presentation.slides) > MAX_PAGES:
        raise ParseFailure("PPTX 幻灯片数量超过安全上限。")
    parts: list[str] = []
    locations: list[dict[str, Any]] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        order = 0
        for shape in slide.shapes:
            for value in _shape_text(shape):
                order += 1
                _append(parts, locations, value, slide=slide_number, shape_order=order)
        notes_slide = getattr(slide, "notes_slide", None)
        if notes_slide is not None:
            for shape in notes_slide.notes_text_frame.paragraphs:
                note = shape.text.strip()
                if note:
                    order += 1
                    _append(parts, locations, note, slide=slide_number, note=True, shape_order=order)
    return {
        "text": "\n".join(parts),
        "metadata": {
            "slide_count": len(presentation.slides),
            "character_count": sum(len(p) for p in parts),
        },
        "locations": locations,
    }


def _parse_ppt(path: Path, executable_value: str) -> dict[str, Any]:
    from pypdf import PdfReader

    executable = Path(executable_value) if executable_value else None
    if executable is None or not executable.is_file():
        raise ParseFailure("缺少本地 LibreOffice，无法解析旧版 PowerPoint 文件。")
    with tempfile.TemporaryDirectory(prefix="mindmate-ppt-parse-") as temporary:
        workdir = Path(temporary)
        source = workdir / "source.ppt"
        source.write_bytes(path.read_bytes())
        profile = workdir / "profile"
        profile.mkdir()
        command = [
            str(executable),
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--norestore",
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--convert-to",
            "pdf:impress_pdf_Export",
            "--outdir",
            str(workdir),
            str(source),
        ]
        env = {
            key: os.environ[key]
            for key in ("SystemRoot", "WINDIR", "PATH", "USERPROFILE", "APPDATA", "LOCALAPPDATA")
            if os.environ.get(key)
        }
        env["TEMP"] = str(workdir)
        env["TMP"] = str(workdir)
        process: subprocess.Popen[bytes] | None = None
        try:
            process = subprocess.Popen(
                command,
                cwd=workdir,
                stdin=subprocess.DEVNULL,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                env=env,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
            process.wait(timeout=25)
        except subprocess.TimeoutExpired as exc:
            if process is not None:
                process.kill()
                process.wait()
            raise ParseFailure("旧版 PowerPoint 转换超时。") from exc
        except OSError as exc:
            raise ParseFailure("无法启动旧版 PowerPoint 转换器。") from exc
        pdf_path = workdir / "source.pdf"
        if process is None or process.returncode != 0 or not pdf_path.is_file():
            raise ParseFailure("旧版 PowerPoint 无法转换为可解析页面。")
        if pdf_path.stat().st_size > 200 * 1024 * 1024:
            raise ParseFailure("旧版 PowerPoint 转换结果超过安全上限。")
        reader = PdfReader(pdf_path, strict=True)
        if reader.is_encrypted or len(reader.pages) > MAX_PAGES:
            raise ParseFailure("旧版 PowerPoint 页数超过安全上限或文件已加密。")
        parts: list[str] = []
        locations: list[dict[str, Any]] = []
        for page_number, page in enumerate(reader.pages, start=1):
            _append(parts, locations, page.extract_text() or "", slide=page_number)
        return {
            "text": "\n".join(parts),
            "metadata": {"slide_count": len(reader.pages), "character_count": sum(map(len, parts))},
            "locations": locations,
        }


def _parse_xlsx(path: Path) -> dict[str, Any]:
    def local_name(tag: str) -> str:
        return tag.rsplit("}", 1)[-1]

    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            names = [item.filename.replace("\\", "/") for item in entries]
            if len(entries) > MAX_ARCHIVE_ENTRIES or len(names) != len(set(names)):
                raise ParseFailure("Excel 文件包条目数量超限或存在重复条目。")
            expanded = sum(item.file_size for item in entries)
            compressed = sum(max(1, item.compress_size) for item in entries)
            if (
                expanded > MAX_ARCHIVE_EXPANDED_SIZE
                or expanded > compressed * MAX_ARCHIVE_COMPRESSION_RATIO
            ):
                raise ParseFailure("Excel 文件包展开体积或压缩比例超过安全上限。")
            name_set = set(names)
            shared_strings: list[str] = []
            if "xl/sharedStrings.xml" in name_set:
                root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
                for item in root:
                    if local_name(item.tag) != "si":
                        continue
                    shared_strings.append(
                        "".join(
                            text.text or "" for text in item.iter() if local_name(text.tag) == "t"
                        )
                    )
                    if len(shared_strings) > MAX_SPREADSHEET_SHARED_STRINGS:
                        raise ParseFailure("Excel 共享字符串数量超过安全上限。")
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            relationships: dict[str, str] = {}
            if "xl/_rels/workbook.xml.rels" in name_set:
                rel_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
                for relation in rel_root:
                    if local_name(relation.tag) == "Relationship":
                        target = relation.attrib.get("Target", "")
                        if target.startswith("/") or ".." in Path(target).parts:
                            continue
                        relationships[relation.attrib.get("Id", "")] = "xl/" + target.lstrip("/")
            sheets = []
            for sheet in workbook.iter():
                if local_name(sheet.tag) == "sheet":
                    rel_id = next(
                        (value for key, value in sheet.attrib.items() if key.endswith("}id")), ""
                    )
                    target = relationships.get(rel_id)
                    if target in name_set:
                        sheets.append((sheet.attrib.get("name", "工作表"), target))
            parts: list[str] = []
            locations: list[dict[str, Any]] = []
            cell_count = 0
            for sheet_index, (sheet_name, target) in enumerate(sheets, start=1):
                root = ET.fromstring(archive.read(target))
                for cell in (node for node in root.iter() if local_name(node.tag) == "c"):
                    cell_count += 1
                    if cell_count > MAX_SPREADSHEET_CELLS:
                        raise ParseFailure("Excel 单元格数量超过安全上限。")
                    address = cell.attrib.get("r", "")
                    cell_type = cell.attrib.get("t", "")
                    value_node = next((node for node in cell if local_name(node.tag) == "v"), None)
                    if cell_type == "inlineStr":
                        value = "".join(
                            node.text or "" for node in cell.iter() if local_name(node.tag) == "t"
                        )
                    elif value_node is None:
                        continue
                    elif cell_type == "s":
                        try:
                            value = shared_strings[int(value_node.text or "")]
                        except (ValueError, IndexError):
                            continue
                    else:
                        value = value_node.text or ""
                    _append(
                        parts,
                        locations,
                        value,
                        sheet=sheet_index,
                        sheet_name=sheet_name,
                        cell=address,
                    )
            if not sheets:
                raise ParseFailure("Excel 工作簿没有可读取的工作表。")
            text = "\n".join(parts)
            return {
                "text": text,
                "metadata": {
                    "sheet_count": len(sheets),
                    "xlsx_sheets": [name for name, _ in sheets],
                    "cell_count": cell_count,
                    "character_count": len(text),
                },
                "locations": locations,
            }
    except (OSError, zipfile.BadZipFile, ET.ParseError, KeyError, RuntimeError) as exc:
        raise ParseFailure("Excel 文件包损坏或内容无法读取。") from exc


def _title_text(value: Any) -> str:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, dict):
        for key in ("plainText", "text", "title"):
            candidate = value.get(key)
            if isinstance(candidate, str) and candidate.strip():
                return candidate.strip()
    return ""


def _parse_xmind_json(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        raw_sheets = payload
    elif isinstance(payload, dict):
        raw_sheets = payload.get("sheets")
        if not isinstance(raw_sheets, list):
            workbook = payload.get("workbook")
            raw_sheets = workbook.get("sheets") if isinstance(workbook, dict) else None
        if not isinstance(raw_sheets, list):
            raw_sheets = [payload] if isinstance(payload.get("rootTopic"), dict) else []
    else:
        raw_sheets = []

    sheets: list[dict[str, Any]] = []
    for index, raw_sheet in enumerate(raw_sheets[:MAX_PAGES], start=1):
        if not isinstance(raw_sheet, dict):
            continue
        root = raw_sheet.get("rootTopic")
        if not isinstance(root, dict):
            continue
        sheets.append(
            {
                "title": _title_text(raw_sheet.get("title")) or f"工作表 {index}",
                "root": root,
            }
        )
    return sheets


def _parse_xmind_xml(payload: bytes) -> list[dict[str, Any]]:
    try:
        root = ET.fromstring(payload)
    except ET.ParseError as exc:
        raise ParseFailure("XMind XML 结构无效。") from exc

    def local_name(tag: str) -> str:
        return tag.rsplit("}", 1)[-1]

    def child(element: ET.Element, name: str) -> ET.Element | None:
        return next((item for item in element if local_name(item.tag) == name), None)

    sheets: list[dict[str, Any]] = []
    for index, sheet in enumerate((item for item in root.iter() if local_name(item.tag) == "sheet"), start=1):
        topic = child(sheet, "topic")
        if topic is None:
            continue

        def convert_topic(element: ET.Element, depth: int = 0) -> dict[str, Any]:
            if depth > MAX_XMIND_DEPTH:
                raise ParseFailure("XMind 主题层级超过安全上限。")
            title_element = child(element, "title")
            title = "".join(title_element.itertext()).strip() if title_element is not None else ""
            children: list[dict[str, Any]] = []
            container = child(element, "children")
            if container is not None:
                for topics in container:
                    if local_name(topics.tag) != "topics":
                        continue
                    for nested in topics:
                        if local_name(nested.tag) == "topic":
                            children.append(convert_topic(nested, depth + 1))
            return {"id": element.get("id"), "title": title, "children": children}

        sheet_title = child(sheet, "title")
        sheets.append(
            {
                "title": "".join(sheet_title.itertext()).strip() if sheet_title is not None else f"工作表 {index}",
                "root": convert_topic(topic),
            }
        )
        if len(sheets) >= MAX_PAGES:
            break
    return sheets


def _parse_xmind(path: Path) -> dict[str, Any]:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            names_in_order = [entry.filename.replace("\\", "/") for entry in entries]
            if len(names_in_order) != len(set(names_in_order)):
                raise ParseFailure("XMind 文件包含重复 ZIP 条目。")
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise ParseFailure("XMind 文件条目数量超过安全上限。")
            expanded_size = sum(entry.file_size for entry in entries)
            compressed_size = sum(max(1, entry.compress_size) for entry in entries)
            if expanded_size > MAX_ARCHIVE_EXPANDED_SIZE:
                raise ParseFailure("XMind 解压后体积超过安全上限。")
            if expanded_size > compressed_size * MAX_ARCHIVE_COMPRESSION_RATIO:
                raise ParseFailure("XMind 压缩比例超过安全上限。")
            names = {entry.filename.replace("\\", "/") for entry in entries}
            json_name = "content.json" if "content.json" in names else None
            xml_name = "content.xml" if "content.xml" in names else None
            if json_name is not None:
                content_info = archive.getinfo(json_name)
                if content_info.file_size > MAX_XMIND_CONTENT_SIZE:
                    raise ParseFailure("XMind 主数据超过安全解析上限。")
                sheets = _parse_xmind_json(json.loads(archive.read(content_info).decode("utf-8-sig")))
            elif xml_name is not None:
                content_info = archive.getinfo(xml_name)
                if content_info.file_size > MAX_XMIND_CONTENT_SIZE:
                    raise ParseFailure("XMind 主数据超过安全解析上限。")
                sheets = _parse_xmind_xml(archive.read(content_info))
            else:
                raise ParseFailure("XMind 文件缺少受支持的 content.json 或 content.xml。")
    except (OSError, zipfile.BadZipFile, RuntimeError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ParseFailure("XMind 文件包损坏或内容无法读取。") from exc

    parts: list[str] = []
    locations: list[dict[str, Any]] = []
    outline: list[dict[str, Any]] = []
    topic_count = 0

    def visit_topic(
        topic: dict[str, Any],
        sheet_index: int,
        parent_id: str | None,
        depth: int,
    ) -> dict[str, Any] | None:
        nonlocal topic_count
        if depth > MAX_XMIND_DEPTH:
            raise ParseFailure("XMind 主题层级超过安全上限。")
        topic_count += 1
        if topic_count > MAX_XMIND_TOPICS:
            raise ParseFailure("XMind 主题数量超过安全上限。")
        title = _title_text(topic.get("title"))
        topic_id = topic.get("id") if isinstance(topic.get("id"), str) else None
        node = {"title": title, "children": []}
        if title:
            _append(parts, locations, title, sheet=sheet_index, topic_id=topic_id, parent_id=parent_id, depth=depth)
        children_value = topic.get("children")
        child_topics: list[Any] = []
        if isinstance(children_value, dict):
            for key in ("attached", "detached"):
                value = children_value.get(key)
                if isinstance(value, list):
                    child_topics.extend(value)
        if isinstance(children_value, list):
            child_topics.extend(children_value)
        for child_topic in child_topics:
            if isinstance(child_topic, dict):
                child_node = visit_topic(child_topic, sheet_index, topic_id, depth + 1)
                if child_node is not None:
                    node["children"].append(child_node)
        return node if title or node["children"] else None

    for sheet_index, sheet in enumerate(sheets, start=1):
        sheet_title = str(sheet["title"]).strip() or f"工作表 {sheet_index}"
        _append(parts, locations, sheet_title, sheet=sheet_index, kind="sheet")
        root_node = visit_topic(sheet["root"], sheet_index, None, 0)
        outline.append({"title": sheet_title, "root": root_node})
    if topic_count == 0:
        raise ParseFailure("XMind 文件中没有可读取的主题。")
    text = "\n".join(parts)
    return {
        "text": text,
        "metadata": {
            "sheet_count": len(outline),
            "topic_count": topic_count,
            "character_count": len(text),
            "xmind_outline": outline,
        },
        "locations": locations,
    }


def main() -> int:
    if len(sys.argv) != 5:
        return 2
    document_type, source_value, output_value, libreoffice_value = sys.argv[1:]
    source = Path(source_value).resolve(strict=True)
    output = Path(output_value).resolve()
    parsers = {"PDF": _parse_pdf, "DOCX": _parse_docx, "PPT": _parse_ppt, "PPTX": _parse_pptx, "XLSX": _parse_xlsx, "XMIND": _parse_xmind}
    parser = parsers.get(document_type)
    if parser is None:
        raise ParseFailure("解析类型不受支持。")
    payload = _parse_ppt(source, libreoffice_value) if document_type == "PPT" else parser(source)
    output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(str(exc)[:500], file=sys.stderr)
        raise SystemExit(1) from None
